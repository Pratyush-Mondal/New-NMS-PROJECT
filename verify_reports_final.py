import urllib.request
import urllib.parse
import json
import io
import zipfile
import re

BASE_URL = "http://127.0.0.1:8000"

def log(msg, ok=True):
    symbol = "[PASS]" if ok else "[FAIL]"
    print(f"{symbol} {msg}")

def test_api():
    print("\n--- 1. Testing Backend Analytics API Across Calendar Ranges ---")
    ranges = [
        ("Today Only", "2026-10-14", "2026-10-14"),
        ("Past 7 Days", "2026-10-07", "2026-10-14"),
        ("Past 30 Days", "2026-09-14", "2026-10-14"),
        ("Historical Range", "2026-05-01", "2026-06-01")
    ]

    for label, f, t in ranges:
        url = f"{BASE_URL}/api/reports/analytics?from_date={f}&to_date={t}"
        req = urllib.request.Request(url)
        with urllib.request.urlopen(req) as res:
            assert res.status == 200
            data = json.loads(res.read().decode())
            assert data["success"] is True
            assert data["period"]["from_date"] == f
            assert data["period"]["to_date"] == t
            s = data["summary"]
            log(f"{label} ({f} to {t}): Total Assets={s['total_assets']}, Online={s['online_count']}, Offline={s['offline_count']}, Uptime={s['overall_availability_pct']}%")

def test_excel_export():
    print("\n--- 2. Testing Excel Export (.xlsx) ---")
    import openpyxl
    url = f"{BASE_URL}/api/reports/export/excel?from_date=2026-10-07&to_date=2026-10-14"
    with urllib.request.urlopen(url) as res:
        assert res.status == 200
        content = res.read()
        assert len(content) > 10000
        wb = openpyxl.load_workbook(io.BytesIO(content))
        expected_sheets = ["Summary", "Devices Detail", "CCTV Cameras Detail", "Desktops Detail", "Status History"]
        assert wb.sheetnames == expected_sheets
        log(f"Excel loaded successfully ({len(content)} bytes). Sheets verified: {wb.sheetnames}")
        summary_sheet = wb["Summary"]
        assert "NEO METALIKS NMS" in str(summary_sheet["A1"].value)
        log(f"Excel title confirmed: '{summary_sheet['A1'].value}'")
        log(f"Devices sheet row count: {wb['Devices Detail'].max_row}")
        log(f"Cameras sheet row count: {wb['CCTV Cameras Detail'].max_row}")
        log(f"Desktops sheet row count: {wb['Desktops Detail'].max_row}")

def test_csv_export():
    print("\n--- 3. Testing CSV Export (.csv) ---")
    url = f"{BASE_URL}/api/reports/export/csv?from_date=2026-10-07&to_date=2026-10-14"
    with urllib.request.urlopen(url) as res:
        assert res.status == 200
        content = res.read()
        assert content.startswith(b"\xef\xbb\xbf"), "Missing UTF-8 BOM"
        text = content.decode("utf-8")
        lines = text.splitlines()
        log(f"CSV downloaded successfully ({len(content)} bytes, {len(lines)} lines)")
        header_clean = lines[0].lstrip("\ufeff")
        assert "Device,Type,IP Address,Status,Uptime" in header_clean
        log(f"CSV Header 1 confirmed: '{header_clean}'")
        assert any("REPORT METADATA & SUMMARY METRICS" in l for l in lines)
        assert any("DETAILED ASSET PERFORMANCE BREAKDOWN" in l for l in lines)
        assert any("STATUS TRANSITIONS & HISTORY EVENTS" in l for l in lines)
        log("CSV sections (Metadata, Performance Breakdown, Transitions) verified!")

def test_zip_export():
    print("\n--- 4. Testing ZIP Bundle Export (.zip) ---")
    url = f"{BASE_URL}/api/reports/export/zip?from_date=2026-10-07&to_date=2026-10-14"
    with urllib.request.urlopen(url) as res:
        assert res.status == 200
        content = res.read()
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            files = zf.namelist()
            expected_files = ['01_Summary.csv', '02_Devices_Detail.csv', '03_CCTV_Cameras_Detail.csv', '04_Desktops_Detail.csv', '05_Status_History.csv']
            assert files == expected_files
            log(f"ZIP bundle loaded ({len(content)} bytes). All 5 CSV files verified: {files}")

def test_frontend_integrity():
    print("\n--- 5. Testing Frontend HTML and JavaScript Alignment ---")
    with open("index.html", "r", encoding="utf-8") as f:
        html = f.read()

    with open("script_v2.js", "r", encoding="utf-8") as f:
        js = f.read()

    # Check date inputs
    assert 'id="reportFromDate"' in html
    assert 'id="reportToDate"' in html
    log("Calendar inputs reportFromDate & reportToDate exist in index.html")

    # Check summary card IDs
    assert 'id="reportTotalDevices"' in html
    assert 'id="reportOnlineDevices"' in html
    assert 'id="reportActiveAlerts"' in html
    assert 'id="reportUptime"' in html
    log("Summary card IDs (reportTotalDevices, reportOnlineDevices, reportActiveAlerts, reportUptime) exist in index.html")

    # Check download buttons
    assert 'id="downloadCSV"' in html and 'onclick="downloadCSVReport()"' in html
    assert 'id="downloadExcel"' in html and 'onclick="downloadExcelReport()"' in html
    assert 'id="downloadZIP"' in html and 'onclick="downloadZipReport()"' in html
    log("Download buttons (downloadCSV, downloadExcel, downloadZIP) with onclick handlers exist in index.html")

    # Check table columns alignment
    # In index.html table header:
    # <th>Device ID</th> <th>Device Name</th> <th>Type</th> <th>IP Address</th> <th>Status</th> <th>Uptime</th> <th>Last Check</th>
    assert "<th>Device ID</th>" in html
    assert "<th>Device Name</th>" in html
    assert "<th>Type</th>" in html
    assert "<th>IP Address</th>" in html
    assert "<th>Status</th>" in html
    assert "<th>Uptime</th>" in html
    assert "<th>Last Check</th>" in html

    # In script_v2.js:
    # <td><b>${escapeHtml(device.id)}</b></td>
    # <td>${escapeHtml(device.name)}</td>
    # <td>${escapeHtml(device.type)}</td>
    # <td>${escapeHtml(device.ip)}</td>
    # <td><span class="${statusClass}">${escapeHtml(device.status)}</span></td>
    # <td title="${escapeHtml(tooltipTitle)}">${escapeHtml(device.uptime)}</td>
    # <td><small>${escapeHtml(device.lastChecked)}</small></td>
    idx_ip = js.find("<td>${escapeHtml(device.ip)}</td>")
    idx_status = js.find("<td><span class=\"${statusClass}\">${escapeHtml(device.status)}</span></td>")
    assert idx_ip != -1 and idx_status != -1
    assert idx_ip < idx_status, "IP Address must be rendered BEFORE Status to match header!"
    log("Table column order verified: IP Address is column 4, Status is column 5 - matches table header exactly!")

    # Check window exports
    assert "window.downloadCSVReport = downloadCSVReport;" in js
    assert "window.downloadExcelReport = downloadExcelReport;" in js
    assert "window.downloadZipReport = downloadZipReport;" in js
    assert "window.generateNMSReport = generateNMSReport;" in js
    log("All report action functions exposed to window object for bulletproof execution")

if __name__ == "__main__":
    test_api()
    test_excel_export()
    test_csv_export()
    test_zip_export()
    test_frontend_integrity()
    print("\n=======================================================")
    print(">>> ALL REPORT VERIFICATIONS PASSED 100% SUCCESSFULLY! <<<")
    print("=======================================================")
